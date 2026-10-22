#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Keyword Gap Radar — 流量词短板与机会雷达
天机OS · amazon-keyword-gap-radar 技能核心模块

输入: 卖家精灵「多产品对比流量词/发现流量词」导出 xlsx
输出: 打标xlsx(汇总+明细) + JSON摘要 + 文本简报

六类标签(互斥, 按优先级判定):
  LEAK   烧而不沉  我方广告份额高但自然≈0 → AOT L3 词级断链
  DEFEND 防守词    我方自然份额高, 护城河
  OPP    临门一脚  我方自然已收录(0.5~2%)未打透
  GAP    短板词    竞品自然验证过, 我方缺席
  BLUE   蓝海词    搜索量可观, 全员份额低
  TAIL   长尾观察  其余

用法:
  python gap_radar.py --input 导出.xlsx --self B0XXXXXXXX [--out-dir .] [--top 30]
阈值配置: assets/weights.json (全部工程假设, 可调)
"""
import argparse, json, math, re, sys, warnings
from pathlib import Path
warnings.filterwarnings("ignore")

DEFAULT_W = {
    "tags": {"leak_ad_min": 1.0, "leak_org_max": 0.5,
             "defend_org_min": 2.0,
             "opp_org_min": 0.5, "opp_org_max": 2.0, "opp_sv_min": 1000,
             "gap_comp_org_min": 2.0, "gap_self_total_max": 0.5,
             "blue_sv_min": 5000, "blue_all_total_max": 1.0},
    "score": {"top3_penalty": True},
    "_source": "engineering_assumption — 阈值为工程假设(份额单位%), 按类目调优"
}


def load_w(path=None):
    p = Path(path) if path else Path(__file__).parent.parent / "assets" / "weights.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else DEFAULT_W


def fnum(v):
    if v in (None, "", "-"):
        return 0.0
    s = str(v).replace("%", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_export(path):
    """解析卖家精灵多产品对比导出: 自动定位表头行与ASIN列。"""
    from openpyxl import load_workbook
    # 注意: 卖家精灵导出的xlsx维度元数据不标准, read_only模式会漏行, 必须普通模式
    ws = load_workbook(path).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = next(i for i, r in enumerate(rows)
                 if r and any(str(c or "").strip() == "关键词" for c in r))
    hdr = [str(c).replace("\n", " ").strip() if c else "" for c in rows[hdr_i]]
    col = {h: i for i, h in enumerate(hdr)}
    asin_cols = {}  # asin -> {total,org,ad}
    for h, i in col.items():
        m = re.match(r"(B0[A-Z0-9]{8})\s+(总|自然|广告)流量份额", h)
        if m:
            asin_cols.setdefault(m.group(1), {})[
                {"总": "total", "自然": "org", "广告": "ad"}[m.group(2)]] = i
    def g(r, name, default=None):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else default
    kws = []
    for r in rows[hdr_i + 1:]:
        if not r or not g(r, "关键词"):
            continue
        kw = {"keyword": str(g(r, "关键词")).strip(),
              "translation": g(r, "翻译") or "",
              "aba_rank": fnum(g(r, "ABA排名")),
              "sv": fnum(g(r, "周搜索量")),
              "bid": fnum(g(r, "建议竞价-推荐")),
              "kw_cvr": fnum(g(r, "关键词点击转化率")),
              "top3_click": fnum(g(r, "ABA Top3集中度-点击")),
              "shares": {}}
        for asin, cc in asin_cols.items():
            kw["shares"][asin] = {k: fnum(r[i]) if i < len(r) else 0.0
                                  for k, i in cc.items()}
        kws.append(kw)
    return kws, sorted(asin_cols.keys())


def tag_one(kw, self_asin, comps, T):
    s = kw["shares"].get(self_asin, {})
    st, so, sa = s.get("total", 0), s.get("org", 0), s.get("ad", 0)
    comp_org_max = max((kw["shares"][c].get("org", 0) for c in comps), default=0)
    all_total_max = max((kw["shares"][a].get("total", 0) for a in kw["shares"]), default=0)
    lg = math.log10(kw["sv"] + 1)
    pen = 1 + kw["top3_click"] / 100.0
    if sa >= T["leak_ad_min"] and so < T["leak_org_max"]:
        return "LEAK 烧而不沉", round(lg * sa, 2), "广告在烧但自然未沉淀(AOT-L3词级): 查该词CVR是否低于市场; 达标则续投等沉淀, 不达标降位/换匹配"
    if so >= T["defend_org_min"]:
        return "DEFEND 防守词", round(lg * so, 2), "护城河词: 保排名, 广告降配, 监控竞品份额侵蚀"
    if T["opp_org_min"] <= so < T["opp_org_max"] and kw["sv"] >= T["opp_sv_min"]:
        return "OPP 临门一脚", round(lg * so * 2 / pen, 2), "已收录未打透: SP集中加码, 按SPR/8定日单量目标冲前3页"
    if comp_org_max >= T["gap_comp_org_min"] and st < T["gap_self_total_max"]:
        return "GAP 短板词", round(lg * comp_org_max / pen, 2), "竞品自然验证过而我方缺席: 先补收录(埋词/ST), 再SP低位测试"
    if kw["sv"] >= T["blue_sv_min"] and all_total_max < T["blue_all_total_max"]:
        return "BLUE 蓝海词", round(lg, 2), "全员份额低的大词: 验证相关性后低价卡位, 差异化埋词"
    return "TAIL 长尾观察", round(lg * max(so, comp_org_max) / 10, 2), "暂不投入, 季度复查"


def analyze(kws, self_asin, comps, W):
    T = W["tags"]
    for kw in kws:
        kw["tag"], kw["score"], kw["action"] = tag_one(kw, self_asin, comps, T)
    order = ["LEAK 烧而不沉", "GAP 短板词", "OPP 临门一脚",
             "DEFEND 防守词", "BLUE 蓝海词", "TAIL 长尾观察"]
    summary = {t: [k for k in kws if k["tag"] == t] for t in order}
    for t in summary:
        summary[t].sort(key=lambda k: k["score"], reverse=True)
    return summary, order


def write_xlsx(kws, summary, order, self_asin, comps, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    FILL = {"LEAK 烧而不沉": "FFC7CE", "GAP 短板词": "FFEB9C", "OPP 临门一脚": "C6EFCE",
            "DEFEND 防守词": "BDD7EE", "BLUE 蓝海词": "D9E1F2", "TAIL 长尾观察": "F2F2F2"}
    wb = Workbook()
    # ---- 明细 sheet ----
    det = wb.active
    det.title = "明细"
    hdr = ["标签", "优先分", "关键词", "翻译", "周搜索量", "ABA排名",
           f"我方自然% ({self_asin})", "我方广告%", "我方总%"]
    for c in comps:
        hdr.append(f"竞品自然% ({c})")
    hdr += ["Top3点击集中%", "词CVR%", "建议竞价", "建议动作"]
    det.append(hdr)
    ordered = [k for t in order for k in summary[t]]
    for k in ordered:
        s = k["shares"].get(self_asin, {})
        row = [k["tag"], k["score"], k["keyword"], k["translation"], k["sv"],
               k["aba_rank"] or "", round(s.get("org", 0), 2),
               round(s.get("ad", 0), 2), round(s.get("total", 0), 2)]
        row += [round(k["shares"][c].get("org", 0), 2) for c in comps]
        row += [k["top3_click"], k["kw_cvr"], k["bid"], k["action"]]
        det.append(row)
    n = len(ordered)
    for i, k in enumerate(ordered, start=2):
        det.cell(i, 1).fill = PatternFill("solid", start_color=FILL[k["tag"]])
    bold = Font(bold=True, name="Arial")
    for j in range(1, len(hdr) + 1):
        c = det.cell(1, j)
        c.font = bold
        c.fill = PatternFill("solid", start_color="1F2A44")
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        det.column_dimensions[get_column_letter(j)].width = \
            26 if j == 3 else (34 if j == len(hdr) else 13)
    det.freeze_panes = "A2"
    det.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{n+1}"
    # ---- 汇总 sheet (公式聚合) ----
    sm = wb.create_sheet("汇总", 0)
    sm["A1"] = f"流量词短板与机会雷达 · 我方 {self_asin} vs 竞品 {', '.join(comps)}"
    sm["A1"].font = Font(bold=True, size=13, name="Arial")
    sm.append([])
    sm.append(["标签", "词数", "词数占比", "搜索量合计", "含义"])
    meaning = {"LEAK 烧而不沉": "广告烧钱自然未沉淀(AOT-L3词级断链)",
               "GAP 短板词": "竞品自然验证过、我方缺席",
               "OPP 临门一脚": "已收录未打透, 最优先进攻",
               "DEFEND 防守词": "自然护城河, 保排名",
               "BLUE 蓝海词": "全员低份额大词",
               "TAIL 长尾观察": "暂不投入"}
    r0 = 4
    for i, t in enumerate(order):
        r = r0 + i
        sm.cell(r, 1, t).fill = PatternFill("solid", start_color=FILL[t])
        sm.cell(r, 2, f'=COUNTIF(明细!A:A,A{r})')
        sm.cell(r, 3, f'=B{r}/SUM(B${r0}:B${r0+len(order)-1})').number_format = "0.0%"
        sm.cell(r, 4, f'=SUMIF(明细!A:A,A{r},明细!E:E)').number_format = "#,##0"
        sm.cell(r, 5, meaning[t])
    tr = r0 + len(order)
    sm.cell(tr, 1, "合计").font = bold
    sm.cell(tr, 2, f"=SUM(B{r0}:B{tr-1})").font = bold
    sm.cell(tr, 4, f"=SUM(D{r0}:D{tr-1})").number_format = "#,##0"
    for w_, col_ in zip([18, 10, 10, 14, 46], "ABCDE"):
        sm.column_dimensions[col_].width = w_
    for j in range(1, 6):
        sm.cell(3, j).font = Font(bold=True, name="Arial")
    for row in sm.iter_rows(min_row=1, max_row=tr):
        for c in row:
            if not c.font or not c.font.name:
                c.font = Font(name="Arial")
    out = Path(out_path)
    wb.save(out)
    return out


def brief(summary, order, self_asin, top=10):
    L = [f"【流量词雷达】我方 {self_asin}"]
    counts = " | ".join(f"{t.split()[1]}{len(summary[t])}" for t in order)
    L.append(counts)
    for t in ["LEAK 烧而不沉", "GAP 短板词", "OPP 临门一脚"]:
        ks = summary[t][:top]
        if ks:
            L.append(f"◆ {t} Top{len(ks)}: " +
                     ", ".join(f"{k['keyword']}(SV{int(k['sv'])})" for k in ks[:5]) + " …")
    return "\n".join(L)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--self", dest="self_asin", required=True, help="我方ASIN")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--weights", default=None)
    ap.add_argument("--top", type=int, default=30, help="JSON摘要每类Top N")
    a = ap.parse_args()
    kws, asins = parse_export(a.input)
    if a.self_asin not in asins:
        sys.exit(f"错误: 文件中的ASIN为 {asins}, 未找到 {a.self_asin}")
    comps = [x for x in asins if x != a.self_asin]
    W = load_w(a.weights)
    summary, order = analyze(kws, a.self_asin, comps, W)
    od = Path(a.out_dir)
    od.mkdir(parents=True, exist_ok=True)
    xp = write_xlsx(kws, summary, order, a.self_asin, comps,
                    od / f"流量词雷达_{a.self_asin}.xlsx")
    js = {t: [{k2: k[k2] for k2 in ("keyword", "sv", "score", "action")}
              for k in summary[t][:a.top]] for t in order}
    (od / f"流量词雷达_{a.self_asin}.json").write_text(
        json.dumps(js, ensure_ascii=False, indent=2), encoding="utf-8")
    print(brief(summary, order, a.self_asin))
    print(f"\n→ {xp}")
